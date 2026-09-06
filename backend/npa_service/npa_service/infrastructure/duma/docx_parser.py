"""Extract readable plain text from an OOXML Word document."""

from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from npa_service.application.errors import NpaDocumentError


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


class DocxParser:
    def parse(self, content: bytes) -> str:
        try:
            document = Document(BytesIO(content))
        except (BadZipFile, KeyError, PackageNotFoundError, ValueError) as error:
            raise NpaDocumentError("bill text is not a readable .docx document") from error

        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        tables = [_table_text(table) for table in document.tables]
        text = "\n".join(value for value in paragraphs + tables if value)
        if not text:
            raise NpaDocumentError("bill Word document contains no readable text")
        return text
