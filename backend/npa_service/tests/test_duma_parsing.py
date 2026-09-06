"""Offline tests for strict URLs, Duma HTML, and Word extraction."""

import subprocess
from io import BytesIO

import pytest
from docx import Document
from npa_service.application.errors import InvalidNpaUrlError
from npa_service.infrastructure.duma.doc_parser import DocParser
from npa_service.infrastructure.duma.docx_parser import DocxParser
from npa_service.infrastructure.duma.page_parser import DumaPageParser
from npa_service.infrastructure.duma.url import normalize_duma_bill_url

BILL_URL = "https://sozd.duma.gov.ru/bill/1286425-8"


def test_normalizes_only_exact_duma_bill_urls() -> None:
    assert normalize_duma_bill_url(f"{BILL_URL}/") == BILL_URL
    invalid = (
        "http://sozd.duma.gov.ru/bill/1286425-8",
        "https://evil.example/bill/1286425-8",
        "https://sozd.duma.gov.ru.evil.example/bill/1286425-8",
        "https://sozd.duma.gov.ru/search?q=1286425-8",
    )
    for url in invalid:
        with pytest.raises(InvalidNpaUrlError):
            normalize_duma_bill_url(url)


def test_parses_furthest_stage_latest_date_and_latest_word_text() -> None:
    page = DumaPageParser().parse(_bill_html(), BILL_URL)

    assert page.bill_number == "1286425-8"
    assert page.title == "О тестовом законе (простое пояснение)"
    assert page.stage_code == "arrh_d11"
    assert page.stage == "Опубликование закона"
    assert page.updated_at.isoformat() == "2026-08-04T23:59:59+03:00"
    assert page.published_at == page.updated_at
    assert page.document_url.endswith("/download/new-text")


def test_accepts_pdf_when_word_text_is_not_published() -> None:
    page = DumaPageParser().parse(_pdf_bill_html(), BILL_URL)

    assert page.document_url.endswith("/download/text-pdf")


def test_extracts_paragraphs_and_tables_from_docx() -> None:
    document = Document()
    document.add_paragraph("Статья 1. Новое правило")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Было"
    table.cell(0, 1).text = "Стало"
    buffer = BytesIO()
    document.save(buffer)

    text = DocxParser().parse(buffer.getvalue())

    assert "Статья 1. Новое правило" in text
    assert "Было | Стало" in text


def test_extracts_text_from_legacy_doc_with_antiword(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_run(*args, **kwargs) -> subprocess.CompletedProcess[str]:
        return subprocess.CompletedProcess(args, 0, "Статья 1. Текст законопроекта", "")

    monkeypatch.setattr(subprocess, "run", fake_run)

    assert (
        DocParser().parse(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy")
        == "Статья 1. Текст законопроекта"
    )


def _bill_html() -> str:
    return """
    <span id="number_oz_id">№ 1286425-8</span>
    <span id="oz_name">О тестовом законе</span>
    <p id="oz_name_comment">(простое пояснение)</p>
    <div id="bh_histras"><div id="oz_stages">
      <div class="root-stage" id="arrh_d1"><div class="ttl"><a class="lnk">Внесение</a></div>
        <div class="oz_event" data-eventdate="2026-07-10T10:00:00">
          <a class="a_event_files" href="/download/old-text"><span class="format-msword"></span>
            <span>Текст внесенного законопроекта</span></a></div></div>
      <div class="root-stage" id="arrh_d4"><div class="ttl"><a class="lnk">Второе чтение</a></div>
        <div class="oz_event" data-eventdate="2026-07-22T20:00:00">
          <a class="a_event_files" href="/download/new-text"><span class="format-msword"></span>
            <span>Текст законопроекта ко второму чтению</span></a></div></div>
      <div class="root-stage" id="arrh_d11"><div class="ttl"><a class="lnk">Опубликование закона</a></div>
        <div class="oz_event" data-eventdate="2026-08-04T23:59:59">закон опубликован</div></div>
    </div></div>
    """


def _pdf_bill_html() -> str:
    return """
    <span id="number_oz_id">№ 1286425-8</span>
    <span id="oz_name">О тестовом законе</span>
    <div id="bh_histras"><div id="oz_stages">
      <div class="root-stage" id="arrh_d4"><div class="ttl"><a class="lnk">Первое чтение</a></div>
        <div class="oz_event" data-eventdate="2026-07-22T20:00:00">
          <a class="a_event_files" href="/download/text-pdf"><span class="format-pdf"></span>
            <span>Текст внесенного законопроекта</span></a></div></div>
    </div></div>
    """
